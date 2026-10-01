"""
Konbit — Pake verifikasyon pewòl pou kontab la
Chemen: backend/scripts/pake_verifikasyon.py

Kreye `pake-verifikasyon-pewol.xlsx` (nan dosye kote ou lanse l la):

  - "Kijan pou itilize l": esplikasyon pou kontab la (kreyòl + franse).
  - "Règ yo": tout to ak baremn sistèm lan sèvi. Kontab la ka chanje yon
    selil epi wè efè a sou tout ka yo.
  - "Ka tès yo": 16 ka. Pou chak ka:
      * kalkil la an FÒMIL EXCEL VIZIB, ki swiv règ yo;
      * chif KONMBIT yo, ki soti DIRÈK nan app/routers/payroll.py
        (compute_deductions, employment_factor...);
      * yon kolòn "KONMBIT − Excel" (dwe 0: sinon kòd la pa swiv règ yo);
      * kolòn VID pou kontab la mete pwòp chif pa l. Diferans yo vin wouj.

Lanse l depi rasin repo a:
    backend\\venv\\Scripts\\python.exe -m pip install openpyxl
    backend\\venv\\Scripts\\python.exe backend\\scripts\\pake_verifikasyon.py

openpyxl pa nan requirements.txt: se yon zouti devlopman, pa sèvè a.
Script la pa touche baz done a.
"""

import os
import sys
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from typing import Optional

OUT = Path.cwd() / "pake-verifikasyon-pewol.xlsx"
BACKEND = Path(__file__).resolve().parent.parent
os.chdir(BACKEND)                       # pou .env la (SECRET_KEY) jwenn
sys.path.insert(0, str(BACKEND))
os.environ.setdefault("MAIL_BACKEND", "memory")

from openpyxl import Workbook                                   # noqa: E402
from openpyxl.formatting.rule import FormulaRule                # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

from app.routers import payroll as P                            # noqa: E402

RULES = "'Règ yo'!"
MONEY = '#,##0.00'
HEAD_FILL = PatternFill("solid", fgColor="0B1F4B")
CALC_FILL = PatternFill("solid", fgColor="EEF4FD")
KONMBIT_FILL = PatternFill("solid", fgColor="FFF7E0")
ACCOUNTANT_FILL = PatternFill("solid", fgColor="E8F6EC")
RED_FILL = PatternFill("solid", fgColor="F8C9C9")
THIN = Side(style="thin", color="CBD5E1")


# ---------------------------------------------------------------------------
# KA TÈS YO
# ---------------------------------------------------------------------------

@dataclass
class Case:
    label: str
    salary: float                 # HTG pa peryòd
    start: date
    end: date
    pay: date
    bonus: float = 0.0            # HTG
    overtime_hours: float = 0.0
    hire: Optional[date] = None
    term: Optional[date] = None


SEP_START, SEP_END = date(2026, 9, 1), date(2026, 9, 30)
OCT_PAY, SEP_PAY = date(2026, 10, 5), date(2026, 9, 30)

CASES = [
    Case("Salè piti: anba plafon CFGDCT, IRI 0", 3_000, SEP_START, SEP_END, OCT_PAY),
    Case("Jis anba plafon CFGDCT (4 999)", 4_999, SEP_START, SEP_END, OCT_PAY),
    Case("Jis nan plafon CFGDCT (5 000)", 5_000, SEP_START, SEP_END, OCT_PAY),
    Case("Jis anba egzanpsyon IRI (5 555)", 5_555, SEP_START, SEP_END, OCT_PAY),
    Case("Jis anlè egzanpsyon IRI (5 600)", 5_600, SEP_START, SEP_END, OCT_PAY),
    Case("Tranch 10% (20 000)", 20_000, SEP_START, SEP_END, OCT_PAY),
    Case("Tranch 15% (25 000)", 25_000, SEP_START, SEP_END, OCT_PAY),
    Case("Tranch 25% (45 000)", 45_000, SEP_START, SEP_END, OCT_PAY),
    Case("Tranch 30% (100 000)", 100_000, SEP_START, SEP_END, OCT_PAY),
    Case("Bonis 10 000, peye 30/09/2026 (retni 10%)", 45_000, SEP_START, SEP_END, SEP_PAY, bonus=10_000),
    Case("Bonis 10 000, peye 05/10/2026 (retni 15%)", 45_000, SEP_START, SEP_END, OCT_PAY, bonus=10_000),
    Case("8 èdtan siplemantè, peye 05/10/2026", 45_000, SEP_START, SEP_END, OCT_PAY, overtime_hours=8),
    Case("Anboche 15/09/2026 (prorata)", 45_000, SEP_START, SEP_END, OCT_PAY, hire=date(2026, 9, 15)),
    Case("Dènye jou 10/09/2026 (prorata)", 45_000, SEP_START, SEP_END, OCT_PAY, term=date(2026, 9, 10)),
    Case("Peye chak kenzèn (22 500)", 22_500, date(2026, 9, 1), date(2026, 9, 15), date(2026, 9, 18)),
    Case("Peye chak semèn (10 000) + 4 èdtan sip.", 10_000, date(2026, 9, 7), date(2026, 9, 13),
         date(2026, 9, 16), overtime_hours=4),
]


def konmbit(case: Case) -> dict:
    """Menm kalkil ak payroll._compute_payslip + _recompute, pou yon salarye."""
    period = SimpleNamespace(start_date=case.start, end_date=case.end, pay_date=case.pay)
    emp = SimpleNamespace(hire_date=case.hire or date(2020, 1, 1), termination_date=case.term)

    periods = P._periods_per_year(period)
    salary = round(case.salary * 100)
    base = int(salary * P.employment_factor(emp, period))
    overtime = P.salaried_overtime_pay(salary, periods, case.overtime_hours * 60)
    bonus = round(case.bonus * 100)

    d = P.compute_deductions(
        salary_gross=base, supplemental_gross=overtime + bonus,
        periods_per_year=periods, pay_date=case.pay,
    )
    gross = base + overtime + bonus
    window = P.employment_window(emp, period)
    return {
        "periods": periods,
        "worked": P._business_days(*window) if window[0] else 0,
        "total": P._business_days(case.start, case.end),
        "net": max(0, gross - sum(d.values())),
        **d,
    }


# ---------------------------------------------------------------------------
# FÈY "RÈG YO"
# ---------------------------------------------------------------------------

def rules_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Règ yo")
    ws["A1"] = "Règ KONMBIT itilize (app/routers/payroll.py) — Règles utilisées"
    ws["A1"].font = Font(bold=True, size=13)

    rows = [
        (3, "ONA (anplwaye) — ONA (part salariale)", P.ONA_RATE, "0.00%"),
        (4, "OFATMA (anplwaye)", P.OFATMA_RATE, "0.00%"),
        (5, "CFGDCT", P.CFGDCT_RATE, "0.00%"),
        (6, "Plafon CFGDCT pa mwa (HTG) — Seuil mensuel CFGDCT", P.CFGDCT_MONTHLY_FLOOR / 100, MONEY),
        (7, "FDU / CAS", P.FDU_CAS_RATE, "0.00%"),
        (8, "Abatman sou salè (atik 92) — Abattement", P.SALARY_ABATEMENT, "0.00%"),
        (9, "Retni sou bonis/èdtan sip. AVAN dat chanjman an", P.SUPPLEMENTAL_TAX_RATE_OLD, "0.00%"),
        (10, "Retni sou bonis/èdtan sip. APRE dat chanjman an", P.SUPPLEMENTAL_TAX_RATE_NEW, "0.00%"),
        (11, "Dat chanjman to bonis la — Date du changement", P.SUPPLEMENTAL_TAX_CHANGE_DATE, "DD/MM/YYYY"),
        (12, "Miltiplikatè èdtan siplemantè", P.OVERTIME_MULTIPLIER, "0.00"),
        (13, "Èdtan pa mwa (to orè ekivalan) — Heures par mois", P.HOURS_PER_PERIOD[12], "0.00"),
        (14, "Èdtan pa kenzèn — Heures par quinzaine", P.HOURS_PER_PERIOD[24], "0.00"),
        (15, "Èdtan pa semèn — Heures par semaine", P.HOURS_PER_PERIOD[52], "0.00"),
    ]
    for row, label, value, fmt in rows:
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=value)
        cell.number_format = fmt

    ws["A16"], ws["B16"], ws["C16"] = "Baremn IRI anyèl — De (HTG)", "Jiska / Jusqu'à (HTG)", "To / Taux"
    for c in ("A16", "B16", "C16"):
        ws[c].font = Font(bold=True)
    lower = 0
    for i, (upper, rate) in enumerate(P.TAX_BRACKETS):
        r = 17 + i
        up = upper / 100 if upper is not None else 1e12
        ws.cell(row=r, column=1, value=lower / 100).number_format = MONEY
        ws.cell(row=r, column=2, value=up).number_format = MONEY
        ws.cell(row=r, column=3, value=rate).number_format = "0.00%"
        lower = upper or lower

    notes = [
        "Kijan IRI a kalkile: (salè baz × 90%) × kantite peryòd pa ane → baremn anyèl → ÷ kantite peryòd.",
        "Calcul de l'IRI : (salaire de base × 90 %) × nombre de périodes par an → barème annuel → ÷ nombre de périodes.",
        "Bonis ak èdtan siplemantè: retni FIKS (pa baremn nan, pa abatman). ONA/OFATMA/CFGDCT/FDU sou tout brit la.",
        "CFGDCT: aplike sèlman si brit la, an ekivalan mansyèl, ≥ plafon an.",
        "Tout kalkil yo koupe nan santim (pa awondi) — tous les montants sont tronqués au centime.",
        "Pa nan ka yo: konje san peye (baz − baz/jou × jou san peye: 22 pa mwa, 11 pa kenzèn, 5 pa semèn), avans (apre enpo).",
        "POU VERIFYE: 173,33 èdtan pa mwa = 40 èdtan pa semèn × 52 ÷ 12. Èske se semèn legal la (Kòd Travay: 48 è?)",
    ]
    for i, text in enumerate(notes):
        ws.cell(row=24 + i, column=1, value=text)
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 12


# ---------------------------------------------------------------------------
# FÈY "KA TÈS YO"
# ---------------------------------------------------------------------------

COLUMNS = [
    # (lèt, tit, gwoup)
    ("A", "#", None), ("B", "Ka / Cas", None),
    ("C", "Salè pa peryòd / Salaire par période", "in"), ("D", "Peryòd pa ane / Périodes par an", "in"),
    ("E", "Jou travay / Jours travaillés", "in"), ("F", "Jou nan peryòd la / Jours ouvrés", "in"),
    ("G", "Bonis", "in"), ("H", "Èdtan sip. / Heures sup.", "in"), ("I", "Dat peman / Date de paie", "in"),
    ("J", "Salè baz / Salaire de base", "calc"), ("K", "Èdtan sip. (HTG)", "calc"), ("L", "Brit / Brut", "calc"),
    ("M", "Baz IRI anyèl / Base IRI annuelle", "calc"), ("N", "IRI", "calc"),
    ("O", "Retni bonis / Retenue primes", "calc"), ("P", "ONA", "calc"), ("Q", "OFATMA", "calc"),
    ("R", "CFGDCT", "calc"), ("S", "FDU/CAS", "calc"), ("T", "Total retni / Total retenues", "calc"),
    ("U", "Net (Excel)", "calc"),
    ("V", "IRI", "konmbit"), ("W", "Retni bonis", "konmbit"), ("X", "ONA", "konmbit"),
    ("Y", "OFATMA", "konmbit"), ("Z", "CFGDCT", "konmbit"), ("AA", "FDU/CAS", "konmbit"),
    ("AB", "Net", "konmbit"), ("AC", "KONMBIT − Excel (max)", "konmbit"),
    ("AD", "IRI dapre kontab la / IRI selon le comptable", "acc"),
    ("AE", "Net dapre kontab la / Net selon le comptable", "acc"),
    ("AF", "Diferans net / Écart net", "acc"), ("AG", "Kòmantè / Commentaire", "acc"),
]
GROUP_FILL = {"calc": CALC_FILL, "konmbit": KONMBIT_FILL, "acc": ACCOUNTANT_FILL}


def bracket_term(row: int, b: int) -> str:
    """Yon tranch baremn: ROUNDDOWN(MAX(0, MIN(baz, jiska) − de) × to; 2)."""
    return (f"ROUNDDOWN(MAX(0,MIN(M{row},{RULES}$B${b})-{RULES}$A${b})*{RULES}$C${b},2)")


def cases_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Ka tès yo")
    ws["A1"] = "KONMBIT — Ka tès pewòl / Cas de test de paie (montants en HTG)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Ble: kalkil Excel (fòmil vizib) · Jòn: chif KONMBIT · Vèt: pou kontab la. "
                "Bleu : calcul Excel · Jaune : résultats KONMBIT · Vert : à remplir par le comptable.")

    for letter, title, group in COLUMNS:
        cell = ws[f"{letter}3"]
        cell.value = title
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[letter].width = 34 if letter == "B" else (28 if letter == "AG" else 14)
    ws.row_dimensions[3].height = 48

    for i, case in enumerate(CASES):
        r = 4 + i
        k = konmbit(case)
        values = {
            "A": i + 1, "B": case.label, "C": case.salary, "D": k["periods"],
            "E": k["worked"], "F": k["total"], "G": case.bonus, "H": case.overtime_hours, "I": case.pay,
            "J": f"=ROUNDDOWN(C{r}*E{r}/F{r},2)",
            "K": (f"=ROUNDDOWN(C{r}/IF(D{r}=12,{RULES}$B$13,IF(D{r}=24,{RULES}$B$14,{RULES}$B$15))"
                  f"*H{r}*{RULES}$B$12,2)"),
            "L": f"=J{r}+K{r}+G{r}",
            "M": f"=ROUNDDOWN(J{r}*(1-{RULES}$B$8),2)*D{r}",
            "N": "=ROUNDDOWN((" + "+".join(bracket_term(r, b) for b in range(17, 22)) + f")/D{r},2)",
            "O": f"=ROUNDDOWN((K{r}+G{r})*IF(I{r}>={RULES}$B$11,{RULES}$B$10,{RULES}$B$9),2)",
            "P": f"=ROUNDDOWN(L{r}*{RULES}$B$3,2)",
            "Q": f"=ROUNDDOWN(L{r}*{RULES}$B$4,2)",
            "R": f"=IF(ROUNDDOWN(L{r}*D{r}/12,2)>={RULES}$B$6,ROUNDDOWN(L{r}*{RULES}$B$5,2),0)",
            "S": f"=ROUNDDOWN(L{r}*{RULES}$B$7,2)",
            "T": f"=SUM(N{r}:S{r})",
            "U": f"=MAX(0,L{r}-T{r})",
            "V": k["tax_amount"] / 100, "W": k["supplemental_tax_amount"] / 100,
            "X": k["ona_amount"] / 100, "Y": k["ofatma_amount"] / 100,
            "Z": k["cfgdct_amount"] / 100, "AA": k["fdu_cas_amount"] / 100, "AB": k["net"] / 100,
            "AC": (f"=MAX(ABS(V{r}-N{r}),ABS(W{r}-O{r}),ABS(X{r}-P{r}),ABS(Y{r}-Q{r}),"
                   f"ABS(Z{r}-R{r}),ABS(AA{r}-S{r}),ABS(AB{r}-U{r}))"),
            "AF": f'=IF(AE{r}="","",ROUND(AB{r}-AE{r},2))',
        }
        for letter, _, group in COLUMNS:
            cell = ws[f"{letter}{r}"]
            if letter in values:
                cell.value = values[letter]
            if group in GROUP_FILL:
                cell.fill = GROUP_FILL[group]
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if letter == "I":
                cell.number_format = "DD/MM/YYYY"
            elif letter not in ("A", "B", "D", "E", "F", "H", "AG"):
                cell.number_format = MONEY

    last = 3 + len(CASES)
    ws.conditional_formatting.add(f"AC4:AC{last}", FormulaRule(formula=["AC4>0.015"], fill=RED_FILL))
    ws.conditional_formatting.add(
        f"AF4:AF{last}", FormulaRule(formula=['AND(AF4<>"",ABS(AF4)>0.015)'], fill=RED_FILL))
    ws.freeze_panes = "C4"


def help_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Kijan pou itilize l"
    lines = [
        ("KONMBIT — Pake verifikasyon pewòl", True),
        ("", False),
        ("KREYÒL", True),
        ("1. Fèy « Règ yo » gen tout to ak baremn KONMBIT sèvi. Verifye yo ak lwa a.", False),
        ("2. Fèy « Ka tès yo » gen 16 ka. Kolòn ble yo se kalkil la an fòmil Excel: klike sou yon selil pou wè l.", False),
        ("3. Kolòn jòn yo se chif KONMBIT bay. « KONMBIT − Excel » dwe 0 (sinon li vin wouj).", False),
        ("4. Mete pwòp chif ou nan kolòn vèt yo (IRI ak Net). Diferans lan vin wouj si li pa 0.", False),
        ("5. Si yon règ pa bon, chanje l nan « Règ yo »: tout ka yo rekalkile. Ekri sa w chanje nan Kòmantè.", False),
        ("", False),
        ("FRANÇAIS", True),
        ("1. La feuille « Règ yo » contient tous les taux et le barème utilisés par KONMBIT. Vérifiez-les.", False),
        ("2. La feuille « Ka tès yo » contient 16 cas. Les colonnes bleues sont le calcul en formules Excel visibles.", False),
        ("3. Les colonnes jaunes sont les résultats de KONMBIT. « KONMBIT − Excel » doit être 0 (sinon en rouge).", False),
        ("4. Saisissez vos propres montants dans les colonnes vertes (IRI et Net). L'écart passe en rouge s'il n'est pas nul.", False),
        ("5. Si une règle est fausse, corrigez-la dans « Règ yo » : tous les cas sont recalculés. Notez-le en commentaire.", False),
        ("", False),
        (f"Kreye pa backend/scripts/pake_verifikasyon.py — généré le {date.today():%d/%m/%Y}.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=text).font = Font(bold=bold, size=13 if i == 1 else 11)
    ws.column_dimensions["A"].width = 120


def main() -> int:
    wb = Workbook()
    help_sheet(wb)
    rules_sheet(wb)
    cases_sheet(wb)
    wb.save(OUT)
    print(f"FINI: {OUT}")
    print(f"{len(CASES)} ka tès. Louvri l nan Excel; kolòn « KONMBIT − Excel » dwe 0 tout kote.")
    return 0


if __name__ == "__main__":
    sys.exit(main())