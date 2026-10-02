"""
Tès done biznis la (app/org_data.py, routers/org_data.py) ak efase kont kandida.
"""
import io
import uuid
import zipfile
from datetime import datetime, timedelta, timezone

from app.database import SessionLocal
from app.models import (
    Application, ApplicationStage, Employee, JobPosting, JobStatus, Organization, User,
)
from openpyxl import load_workbook

from app.org_data import HIDDEN_WORDS, purge_closed_organizations


def _org(slug):
    with SessionLocal() as db:
        return db.query(Organization).filter(Organization.slug == slug).first()


def _zip(client, org):
    r = client.post("/api/org-data/export", json={"password": org["password"]}, headers=org["headers"])
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/zip"
    return zipfile.ZipFile(io.BytesIO(r.content))


def _close(client, org):
    name = _org(org["org_slug"]).name
    r = client.post("/api/org-data/close",
                    json={"password": org["password"], "confirm_name": name}, headers=org["headers"])
    assert r.status_code == 200, r.text


# ---------------------------------------------------------------------------
# EKSPÒTASYON
# ---------------------------------------------------------------------------

def test_export_has_only_this_business_and_no_secrets(client, make_org, make_employee):
    a, b = make_org(), make_org()
    make_employee(a["headers"], first_name="Anayiz", city="=1+2", bank_name="BNC",
                  bank_account_number="1234567890", preferred_payment_method="direct_deposit")
    make_employee(b["headers"], first_name="Zebulon")

    zf = _zip(client, a)
    names = zf.namelist()
    assert {"LI-M.txt", "konmbit-done.xlsx", "csv/employees.csv"} <= set(names)
    employees = zf.read("csv/employees.csv").decode("utf-8-sig")
    assert "Anayiz" in employees and "Zebulon" not in employees
    assert "1234567890" in employees            # dechifre: se done biznis la
    assert "'=1+2" in employees                 # pa gen fòmil Excel ki egzekite

    users = zf.read("csv/users.csv").decode("utf-8-sig")
    assert a["email"] in users and b["email"] not in users

    for name in zf.namelist():
        if not name.endswith(".csv"):
            continue
        text = zf.read(name).decode("utf-8-sig")
        assert "enc:v1:" not in text, name
        header = text.splitlines()[0].lower()
        assert not any(word in header for word in HIDDEN_WORDS), (name, header)
    for table in ("email_tokens", "recovery_codes", "auth_attempts"):
        assert f"csv/{table}.csv" not in names

    # Excel la: yon fèy pa tab, menm done yo, "=1+2" se TÈKS (pa yon fòmil).
    book = load_workbook(io.BytesIO(zf.read("konmbit-done.xlsx")))
    assert "employees" in book.sheetnames and "email_tokens" not in book.sheetnames
    sheet = book["employees"]
    header = [c.value for c in sheet[1]]
    rows = [dict(zip(header, (c.value for c in r))) for r in sheet.iter_rows(min_row=2)]
    ana = next(r for r in rows if r["first_name"] == "Anayiz")
    assert ana["city"] == "=1+2" and ana["bank_account_number"] == "1234567890"
    assert sheet.cell(row=rows.index(ana) + 2, column=header.index("city") + 1).data_type == "s"
    assert not any(any(w in str(h).lower() for w in HIDDEN_WORDS) for h in header)


def test_export_needs_admin_and_password(client, org_admin, make_employee_login):
    hr = make_employee_login(org_admin["headers"], login_role="hr")["headers"]
    assert client.post("/api/org-data/export", json={"password": "x"}, headers=hr).status_code == 403
    bad = client.post("/api/org-data/export", json={"password": "MoveModpas2026"},
                      headers=org_admin["headers"])
    assert bad.status_code == 400


# ---------------------------------------------------------------------------
# FÈMTI AK EFASMAN
# ---------------------------------------------------------------------------

def test_close_needs_exact_name_and_blocks_everyone(client, make_org, make_employee_login):
    org = make_org()
    worker = make_employee_login(org["headers"])["headers"]

    wrong = client.post("/api/org-data/close",
                        json={"password": org["password"], "confirm_name": "pa bon non"},
                        headers=org["headers"])
    assert wrong.status_code == 400

    _close(client, org)
    for headers in (org["headers"], worker):
        assert client.get("/api/auth/identity", headers=headers).status_code in (401, 403)
    assert client.get(f"/api/jobs/public/{org['org_slug']}").status_code == 404


def test_purge_waits_30_days_and_touches_only_that_business(client, make_org, make_employee):
    a, b = make_org(), make_org()
    make_employee(a["headers"])
    make_employee(b["headers"])
    _close(client, a)

    with SessionLocal() as db:
        a_id, b_id = _org(a["org_slug"]).id, _org(b["org_slug"]).id
        assert purge_closed_organizations(db) == []          # poko 30 jou

        org_a = db.query(Organization).filter(Organization.id == a_id).first()
        org_a.closure_requested_at = datetime.now(timezone.utc) - timedelta(days=31)
        db.commit()

        done = purge_closed_organizations(db, apply=True)
        assert [d["organization_id"] for d in done] == [a_id]

    with SessionLocal() as db:
        assert db.query(Organization).filter(Organization.id == a_id).first() is None
        assert db.query(Employee).filter(Employee.organization_id == a_id).count() == 0
        assert db.query(User).filter(User.organization_id == a_id).count() == 0
        assert db.query(Employee).filter(Employee.organization_id == b_id).count() == 1
        assert db.query(User).filter(User.email == b["email"]).count() == 1


# ---------------------------------------------------------------------------
# EFASE KONT KANDIDA
# ---------------------------------------------------------------------------

def test_candidate_deletes_account_but_business_keeps_application(client, org_admin):
    email = f"efase-{uuid.uuid4().hex[:8]}@konbit-test.ht"
    password = "KandidaModpas2026"
    assert client.post("/api/candidate/signup", json={
        "full_name": "Kandida Efase", "email": email, "password": password,
    }).status_code == 200
    login = client.post("/api/auth/login", json={"email": email, "password": password}).json()
    h = {"Authorization": f"Bearer {login['access_token']}"}

    with SessionLocal() as db:
        user_id = db.query(User).filter(User.email == email).first().id
        org = db.query(Organization).filter(Organization.slug == org_admin["org_slug"]).first()
        job = JobPosting(organization_id=org.id, title="Pòs efase", slug=f"efase-{uuid.uuid4().hex[:6]}",
                         status=JobStatus.PUBLISHED, openings=1)
        db.add(job)
        db.flush()
        app = Application(organization_id=org.id, job_posting_id=job.id, full_name="Kandida Efase",
                          email=email, stage=ApplicationStage.RECEIVED, applicant_user_id=user_id)
        db.add(app)
        db.commit()
        app_id = app.id

    assert client.post("/api/candidate/delete-account", json={"password": "MoveModpas2026"},
                       headers=h).status_code == 400
    r = client.post("/api/candidate/delete-account", json={"password": password}, headers=h)
    assert r.status_code == 200, r.text

    assert client.post("/api/auth/login", json={"email": email, "password": password}).status_code == 401
    with SessionLocal() as db:
        app = db.query(Application).filter(Application.id == app_id).first()
        assert app is not None and app.applicant_user_id is None
        assert db.query(User).filter(User.email == email).first() is None

    # Yon kont biznis pa ka sèvi ak bouton sa a
    assert client.post("/api/candidate/delete-account", json={"password": org_admin["password"]},
                       headers=org_admin["headers"]).status_code == 403


def test_closing_kills_every_session_including_refresh(client, make_org):
    org = make_org()
    login = client.post("/api/auth/login", json={"email": org["email"], "password": org["password"]})
    refresh_token = login.json()["refresh_token"]

    _close(client, org)
    again = client.post("/api/auth/refresh", json={"refresh_token": refresh_token})
    assert again.status_code != 200, again.text
